import os
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

EXPORT_FOLDER = "exports"


def _write_sheet(sheet, orders):
    if not orders:
        sheet["A1"] = "No orders were found for the requested timeline."
        sheet["A1"].font = Font(name="Arial")
        return

    headers = list(orders[0].keys())

    for col_num, header in enumerate(headers, start=1):
        cell = sheet.cell(row=1, column=col_num, value=header.replace("_", " ").title())
        cell.font = Font(name="Arial", bold=True)

    for row_num, order in enumerate(orders, start=2):
        for col_num, key in enumerate(headers, start=1):
            value = order.get(key, "")
            if isinstance(value, list):
                value = ", ".join(str(v) for v in value)
            cell = sheet.cell(row=row_num, column=col_num, value=value)
            cell.font = Font(name="Arial")

    for col_num, header in enumerate(headers, start=1):
        max_len = len(header)
        for order in orders:
            value = order.get(header, "")
            if isinstance(value, list):
                value = ", ".join(str(v) for v in value)
            max_len = max(max_len, len(str(value)))
        col_letter = get_column_letter(col_num)
        sheet.column_dimensions[col_letter].width = max_len + 4


def _build_filepath(filepath, label):
    os.makedirs(EXPORT_FOLDER, exist_ok=True)
    if filepath is None:
        timestamp = datetime.now().strftime("%Y-%m-%d_%H%M%S_%f")
        suffix = f"_{label}" if label else ""
        filepath = f"orders_export{suffix}_{timestamp}.xlsx"
    return os.path.join(EXPORT_FOLDER, filepath)


def export_excel_file(orders, filepath=None, label=None):

    filepath = _build_filepath(filepath, label)

    wb = Workbook()
    sheet = wb.active
    sheet.title = "Orders"
    _write_sheet(sheet, orders)

    wb.save(filepath)
    return filepath


def export_excel_file_multi_sheet(sheets, filepath=None, label=None):

    filepath = _build_filepath(filepath, label)

    wb = Workbook()
    first = True

    for sheet_name, orders in sheets.items():
        safe_name = sheet_name[:31]  
        if first:
            sheet = wb.active
            sheet.title = safe_name
            first = False
        else:
            sheet = wb.create_sheet(title=safe_name)
        _write_sheet(sheet, orders)

    wb.save(filepath)
    return filepath